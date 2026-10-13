"""
Extrai o item 4 (Obras de ampliação da capacidade e outras melhorias) da aba PER do modelo econômico-financeiro
da ECO050 (fonte_id eco050_evtea_modelo) e grava, por ano de concessão, os valores de cada linha.

Uso (depois de valores_obras/organizar.py ter baixado a fonte):
    python per_obras/extrair_modelo_eco050.py

Saída em dados/per_obras/:
    modelo_item4_eco050_anual.csv    1 linha por item (folha ou subtotal) x ano de concessão (1..30; ano 1 = F+G)
    modelo_item4_eco050_grupos.csv   1 linha por folha x bloco (valor do bloco e soma de conferência)
Linhas com coluna A = 'INV' são folhas; as demais são subtotais. Folhas sem código na planilha ficam como
`sem_codigo_L<linha>`. Valor do modelo: data-base não informada na planilha (contrato: mai/2012).

4.3 (Elaboração de projetos): nos anos 20-22 o valor fica acima do valor-base constante dos anos 6-30. O excedente
é o projeto das faixas adicionais (4.2.2, anos 20-22; o excedente é 5% do 4.2.2 do mesmo ano) e vai para o bloco
`condicionada`; o valor-base desses anos continua em `manutencao_nivel_servico`. A célula é a mesma nas duas linhas
(coluna `parcela` diz qual parte é).
"""
import sys
import warnings
from pathlib import Path

import openpyxl
import pandas as pd

AQUI = Path(__file__).resolve().parent
PROJETO = AQUI.parent
sys.path.insert(0, str(PROJETO / "valores_obras"))
from fontes import FONTES  # noqa: E402

warnings.filterwarnings("ignore")

FONTE_ID = "eco050_evtea_modelo"
ARQ = PROJETO / "dados" / "valores_obras" / "fontes" / FONTE_ID / "modelo-economico-financeiro-eco-050.xlsx"
SAIDA = PROJETO / "dados" / "per_obras"
DATA_BASE = "não informada (contrato: mai/2012)"
L_INI, L_FIM = 44, 102
COL_TOTAL, COL_AK = 5, 37  # E e AK
COL_ANO1 = (6, 7)  # F, G
COL_ANO2 = 8  # H..AJ = anos 2..30
LETRAS = openpyxl.utils.get_column_letter


L_PROJETOS = 101  # 4.3 ELABORAÇÃO DE PROJETOS
ANOS_PROJ_FAIXAS = (20, 21, 22)  # anos com valor no 4.2.2


def bloco_do(item, linha, ano):
    if linha >= 89 and linha <= 100:  # 4.2 Obras condicionadas (faixas adicionais)
        return "condicionada"
    return "obrigatoria" if ano <= 5 else "manutencao_nivel_servico"


def num(v):
    return float(v) if isinstance(v, (int, float)) else 0.0


def main():
    if not ARQ.exists():
        sys.exit(f"Fonte não encontrada: {ARQ} (rode valores_obras/organizar.py para baixá-la)")
    ws = openpyxl.load_workbook(ARQ, data_only=True)["PER"]
    anual, grupos = [], []
    resumo = []  # (item, nivel, E, soma anual, AK)
    for r in range(L_INI, L_FIM + 1):
        folha = ws.cell(r, 1).value == "INV"
        c = ws.cell(r, 3).value
        item = str(c).strip() if c not in (None, 0, "") else f"sem_codigo_L{r}"
        desc = ws.cell(r, 4).value
        desc = str(desc).strip() if desc not in (None, 0, "") else ""
        nivel = "folha" if folha else "subtotal"
        n0 = len(anual)
        base43 = None
        if r == L_PROJETOS:  # valor-base = valor mais frequente nos anos 6-30
            vals = [round(num(ws.cell(r, COL_ANO2 + a - 2).value), 2) for a in range(6, 31)]
            base43 = max(set(vals), key=vals.count)
        for ano in range(1, 31):
            if ano == 1:
                cels = [f"{LETRAS(c)}{r}" for c in COL_ANO1]
                val = sum(num(ws.cell(r, c).value) for c in COL_ANO1)
                cel = "PER!" + "+".join(cels)
            else:
                col = COL_ANO2 + ano - 2
                val = num(ws.cell(r, col).value)
                cel = f"PER!{LETRAS(col)}{r}"
            comum = dict(item=item, descricao=desc, nivel=nivel, ano_concessao=ano, celula=cel, fonte_id=FONTE_ID,
                         data_base=DATA_BASE, linha_planilha=r)
            if base43 is not None and ano in ANOS_PROJ_FAIXAS and val - base43 > 0.005:
                anual.append(dict(comum, valor=base43, bloco="manutencao_nivel_servico",
                                  parcela="valor-base dos anos 6-30"))
                anual.append(dict(comum, valor=val - base43, bloco="condicionada",
                                  parcela="excedente: projetos das faixas adicionais (4.2.2)"))
            else:
                anual.append(dict(comum, valor=val, bloco=bloco_do(item, r, ano), parcela=""))
        resumo.append((r, item, nivel, num(ws.cell(r, COL_TOTAL).value), sum(a["valor"] for a in anual[n0:]),
                       num(ws.cell(r, COL_AK).value)))

    df = pd.DataFrame(anual)
    g = (df[df.nivel == "folha"].groupby(["linha_planilha", "item", "descricao", "bloco"], sort=False)
         .agg(valor_bloco=("valor", "sum"), anos=("ano_concessao", lambda s: f"{s.min()}-{s.max()}"),
              n_celulas=("valor", "size")).reset_index())
    tot = df[df.nivel == "folha"].groupby("linha_planilha").valor.sum()
    g["soma_folha"] = g.linha_planilha.map(tot)
    g["fonte_id"] = FONTE_ID
    g["data_base"] = DATA_BASE
    SAIDA.mkdir(parents=True, exist_ok=True)
    df.to_csv(SAIDA / "modelo_item4_eco050_anual.csv", sep=";", index=False, encoding="utf-8-sig", decimal=".")
    g.to_csv(SAIDA / "modelo_item4_eco050_grupos.csv", sep=";", index=False, encoding="utf-8-sig", decimal=".")

    # ---- conferências
    TOL = 0.01
    folhas = [x for x in resumo if x[2] == "folha"]
    sub4 = next(x for x in resumo if x[0] == 44)
    soma_folhas = sum(x[3] for x in folhas)
    print(f"Linhas lidas: {len(resumo)} ({len(folhas)} folhas, {len(resumo) - len(folhas)} subtotais); "
          f"{len(df)} linhas em anual, {len(g)} em grupos")
    print(f"Subtotal 4 (E44) = {sub4[3]:,.2f}  (esperado 1.145.919.554,57): "
          f"{'ok' if abs(sub4[3] - 1145919554.57) < TOL else 'DIVERGENTE'}")
    print(f"Soma das folhas (coluna E) = {soma_folhas:,.2f} x E44: {'ok' if abs(soma_folhas - sub4[3]) < TOL else 'DIVERGENTE'}")
    soma_anual_folhas = df[df.nivel == "folha"].valor.sum()
    print(f"Soma anual das folhas = {soma_anual_folhas:,.2f} x E44: {'ok' if abs(soma_anual_folhas - sub4[3]) < TOL else 'DIVERGENTE'}")
    ruins = [x for x in resumo if abs(x[3] - x[4]) > TOL or abs(x[3] - x[5]) > TOL]
    print(f"Soma anual = coluna E = coluna AK em {len(resumo) - len(ruins)} de {len(resumo)} linhas")
    for r, item, nivel, e, s, ak in ruins:
        print(f"  DIVERGENTE L{r} {item} ({nivel}): E={e:,.2f} soma anual={s:,.2f} AK={ak:,.2f}")
    print("Folhas com valor:")
    for x in folhas:
        if abs(x[3]) > 0:
            print(f"  L{x[0]} {x[1]}: {x[3]:,.2f}")
    exc = df[(df.linha_planilha == L_PROJETOS) & (df.bloco == "condicionada")].set_index("ano_concessao").valor
    fa = df[df.item == "4.2.2"].set_index("ano_concessao").valor
    print("4.3 excedente nos anos 20-22 (-> condicionada): " + "; ".join(
        f"ano {a}: {v:,.2f} = {v / fa[a]:.2%} do 4.2.2" for a, v in exc.items()) + f"; total {exc.sum():,.2f}")
    print("Por bloco (folhas):")
    print(g.groupby("bloco").valor_bloco.sum().map("{:,.2f}".format).to_string())
    print(f"Arquivos em {SAIDA}")


if __name__ == "__main__":
    main()
