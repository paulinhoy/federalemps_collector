"""
Consolida o histórico da tabela "Acompanhamento Físico Mensal" (modelo padrão ANTT) dos empreendimentos de
empreendimentos.json.

Uso:
    python acompanhamento_obras/baixar.py            # 1) baixa/atualiza todos os anos
    python acompanhamento_obras/consolidar.py        # 2) consolida
    python acompanhamento_obras/consolidar.py eco050 # opcional: só alguns (slug_antt)

Lê todas as planilhas de dados/acompanhamento_obras/brutos/<slug>/<ano>/ (soltas e dentro dos ZIPs), encontra a aba
do modelo físico mensal pelo cabeçalho e gera em dados/acompanhamento_obras/consolidado/:

    snapshots.csv           um arquivo publicado = um snapshot (mês de referência, data de publicação, qual vale)
    obras.csv               catálogo de obras (obra_id estável, código PER, descrição, local) com 1ª/última aparição
    marcos_por_obra.csv     1 linha por obra: quando surgiu, 1º avanço, conclusão, última mudança, revisões e um
                            histórico em texto
    obras_por_snapshot.csv  TABELA LIMPA: 1 linha por obra por snapshot, campos da obra repetidos, mês de
                            referência, situação e % previsto/executado acumulado até o mês de referência
    etapas_por_snapshot.csv etapas do eventograma (Rio Minas, 4º ano em diante), ligadas à obra-mãe
    serie_mensal.csv        formato longo: snapshot x obra x mês x (previsto, executado), valores como na fonte
    evolucao.csv            por concessão e mês: obras novas/que saíram, replanejamentos, revisões, situação
    mudancas_por_obra.csv   o que mudou em cada obra de um mês para o seguinte
    consolidado.xlsx        todas acima, menos serie_mensal e etapas_por_snapshot (grandes)
"""
import os
import re
import sys
import unicodedata
import warnings
from datetime import datetime
from difflib import SequenceMatcher
from pathlib import Path

import openpyxl
import pandas as pd

from portal_antt import BRUTOS, CONSOLIDADO as SAIDA, empreendimentos, longo

warnings.filterwarnings("ignore")

EMPS = {e["slug_antt"]: e for e in empreendimentos()}
SIGLA = {s: e["sigla"] for s, e in EMPS.items()}
ID_EMP = {s: e["id_empreendimento"] for s, e in EMPS.items()}

MESES = {
    "janeiro": 1, "jan": 1, "fevereiro": 2, "fev": 2, "marco": 3, "mar": 3, "abril": 4, "abr": 4,
    "maio": 5, "mai": 5, "junho": 6, "jun": 6, "julho": 7, "jul": 7, "agosto": 8, "ago": 8,
    "setembro": 9, "set": 9, "outubro": 10, "out": 10, "novembro": 11, "nov": 11, "dezembro": 12, "dez": 12,
}

# colunas da obra, reconhecidas pelo texto do cabeçalho (primeira ocorrência vale)
CAMPOS = [
    ("id_sigicor", r"^id\b|sigicor"),
    ("item_per", r"^item (do )?per"),
    ("descricao", r"^descri"),
    ("frente", r"^frente|^tipo de obra"),
    ("municipio", r"^munic"),
    ("uf", r"^uf$"),
    ("rodovia", r"^rodovia"),
    ("km_inicial", r"^km inicial"),
    ("km_final", r"^km final"),
    ("data_inicio", r"^data de inicio"),
    ("previsao_termino", r"^previsao de termino"),
    ("comparativo", r"^comparativo"),
    ("evento", r"^pesos? dos eventos"),
    ("anos_anteriores", r"^anos anteriores"),
    ("situacao", r"^situacao ate o mes|^status"),
    ("motivo_atraso", r"^motivo"),
    ("ano_per", r"^ano per"),
    ("observacao", r"^observa"),
]


# ----------------------------------------------------------------------------- utilitários

def sem_acento(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def norm(v) -> str:
    return re.sub(r"\s+", " ", sem_acento(str(v))).strip().lower()


def listar_planilhas(slug: str) -> list[Path]:
    base = BRUTOS / slug
    achados = []
    for root, dirs, files in os.walk(longo(base)):
        if "_versoes" in root:
            continue
        for f in files:
            if f.lower().endswith((".xlsx", ".xlsm")) and not f.startswith("~$"):
                achados.append(Path(root.replace("\\\\?\\", "", 1)) / f)
    return sorted(achados)


def numero(v):
    """Converte 0.12, '12%', '0,00%', '-' etc. em float (fração) ou None."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = str(v).strip().replace(" ", "")
    if t in ("", "-", "–", "n/a", "N/A"):
        return None
    pct = t.endswith("%")
    t = t.rstrip("%").replace(".", "").replace(",", ".") if "," in t else t.rstrip("%")
    try:
        x = float(t)
    except ValueError:
        return None
    return x / 100 if pct else x


def km(v):
    """'298+660' -> 298.66 ; 92.7 -> 92.7 ; '-' -> None"""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = str(v).strip().replace(",", ".")
    m = re.match(r"^(\d+)\s*\+\s*(\d+)", t)
    if m:
        return int(m.group(1)) + int(m.group(2)) / 1000
    try:
        return float(t)
    except ValueError:
        return None


def texto(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    t = re.sub(r"\s+", " ", str(v).replace("_x000d_", " ").replace("_x000D_", " ")).strip()
    return t or None


def data(v):
    if isinstance(v, datetime):
        return v.date().isoformat()
    t = texto(v)
    if t:
        m = re.match(r"(\d{2})/(\d{2})/(\d{4})", t)
        if m:
            return f"{m.group(3)}-{m.group(2)}-{m.group(1)}"
    return None


# ----------------------------------------------------------------------------- mês de referência

def datas_do_nome(nome: str):
    """Do nome do arquivo (ou da pasta do ZIP) tira:
    - data de publicação: 'dd_mm_aaaa' ou 'dd_mm_aa';
    - mês dos dados: nome do mês por extenso/abreviado (+ ano), ou 'mm-aaaa'.
    Devolve (publicacao, mes, mes_tem_ano). Ano de 2 dígitos só vale entre 20 e 39 ("11o-ano" não é 2011)."""
    n = sem_acento(nome).lower()
    pub = None
    m = re.search(r"(?<!\d)(\d{2})[_.-](\d{2})[_.-](20\d{2}|[23]\d)(?!\d)", n)
    if m and 1 <= int(m.group(2)) <= 12 and 1 <= int(m.group(1)) <= 31:
        a = int(m.group(3))
        pub = datetime(a + 2000 if a < 100 else a, int(m.group(2)), int(m.group(1)))
        n = n[:m.start()] + " " + n[m.end():]
    padrao = (r"(?<![a-z])(" + "|".join(sorted(MESES, key=len, reverse=True)) +
              r")(?![a-z])[\W_]{0,3}(20\d{2}|[23]\d(?![\do]))?(?!\d)")
    achados = list(re.finditer(padrao, n))
    if achados:  # "janeiro-fevereiro-e-marco" => o último mês citado
        g = achados[-1]
        anos = [x.group(2) for x in achados if x.group(2)]
        a = anos[-1] if anos else None
        ano = (int(a) + 2000 if len(a) == 2 else int(a)) if a else None
        return pub, (MESES[g.group(1)], ano), bool(a)
    m = re.search(r"(?<!\d)(0[1-9]|1[0-2])[-_](20\d{2})(?!\d)", n)
    if m:
        return pub, (int(m.group(1)), int(m.group(2))), True
    return pub, None, False


def resolver_mes(mes_ano, pub, ano_pasta, mes_dados):
    """(mês, ano|None) -> 'aaaa-mm'. Sem ano: usa o da publicação ou da pasta e volta 1 ano se cair
    depois da publicação ou longe do que os dados indicam (ex.: 'dezembro' publicado em janeiro)."""
    mes, ano = mes_ano
    if ano:
        return f"{ano}-{mes:02d}"
    base = pub.year if pub else ano_pasta
    cands = [f"{base}-{mes:02d}", f"{base - 1}-{mes:02d}"]
    if pub:
        cands = [c for c in cands if c <= pub.strftime("%Y-%m")] or cands
    if mes_dados:
        meses = lambda c: int(c[:4]) * 12 + int(c[5:])
        return min(cands, key=lambda c: abs(meses(c) - meses(mes_dados)))
    return cands[0]


# ----------------------------------------------------------------------------- leitura de uma planilha

def achar_aba(wb):
    for ws in wb.worksheets:
        try:
            linhas = [list(r) for r in ws.iter_rows(min_row=1, max_row=15, max_col=80, values_only=True)]
        except Exception:
            continue
        for i, row in enumerate(linhas):
            cel = [norm(c) for c in row if c not in (None, "")]
            if any(c.startswith("comparativo") for c in cel) and any(re.match(r"item (do )?per", c) for c in cel):
                return ws, i
    return None, None


def ler_planilha(caminho: Path):
    wb = openpyxl.load_workbook(longo(caminho), read_only=True, data_only=True)
    ws, hdr = achar_aba(wb)
    if ws is None:
        wb.close()
        return None
    # largura útil: até a última coluna preenchida do cabeçalho/linha de datas (algumas abas têm 16 mil colunas vazias)
    topo = [list(r) for r in ws.iter_rows(min_row=1, max_row=hdr + 3, max_col=400, values_only=True)]
    cab, sub = topo[hdr], topo[hdr + 1] if hdr + 1 < len(topo) else []
    largura = max([j for j, c in enumerate(cab) if c not in (None, "")] +
                  [j for j, c in enumerate(sub) if c not in (None, "")]) + 1

    col = {}
    for j, c in enumerate(cab[:largura]):
        if c in (None, ""):
            continue
        n = norm(c)
        for nome, rx in CAMPOS:
            if nome not in col and re.search(rx, n):
                col[nome] = j
                break

    # blocos = um por cabeçalho "Ano concessão: Nº (...)"; colunas mensais = datas logo abaixo do cabeçalho
    inicios = [(j, texto(c)) for j, c in enumerate(cab[:largura]) if c and "ano concess" in norm(c)]
    meses = {}  # coluna -> (aaaa-mm, nº do bloco)
    for k in (hdr + 1, hdr + 2):
        if k < len(topo):
            ds = {j: c for j, c in enumerate(topo[k][:largura]) if isinstance(c, datetime)}
            if len(ds) >= 6:
                for j in sorted(ds):
                    b = sum(1 for ini, _ in inicios if ini <= j) or 1
                    meses[j] = (ds[j].strftime("%Y-%m"), b)
                break
    rotulos = {i + 1: t for i, (_, t) in enumerate(inicios)}

    # ponto de partida do acumulado = última coluna "anteriores"/"acumulado" antes do 1º mês. No modelo de 2026 da
    # ECO050 o grupo "Anos anteriores" tem 2 colunas (até o 9º ano e "Acumulado até o 10º Ano"); vale a última.
    if meses:
        j0 = min(meses)
        for j in range((col.get("comparativo") or 0) + 1, j0):
            rot = " ".join(norm(x) for x in (cab[j] if j < len(cab) else None, sub[j] if j < len(sub) else None)
                           if x not in (None, "") and not isinstance(x, datetime))
            if "acumulado" in rot or "anteriores" in rot:
                col["anos_anteriores"] = j

    obras, atual, vazias = [], None, 0
    for row in ws.iter_rows(min_row=hdr + 2, max_row=hdr + 6000, max_col=largura, values_only=True):
        comp = norm(row[col["comparativo"]]) if col.get("comparativo") is not None and row[col["comparativo"]] else ""
        tipo = "previsto" if comp.startswith(("prev", "program")) else "executado" if comp.startswith(("exec", "realiz")) else None
        if tipo is None:
            vazias += 1
            if vazias > 60:
                break
            continue
        vazias = 0
        def g(nome):
            j = col.get(nome)
            return row[j] if j is not None and j < len(row) else None
        if tipo == "previsto" or atual is None:
            atual = {
                "id_sigicor": texto(g("id_sigicor")), "item_per": texto(g("item_per")), "descricao": texto(g("descricao")),
                "frente": texto(g("frente")), "evento": texto(g("evento")), "municipio": texto(g("municipio")),
                "uf": texto(g("uf")), "rodovia": texto(g("rodovia")), "km_inicial": km(g("km_inicial")),
                "km_final": km(g("km_final")), "data_inicio": data(g("data_inicio")),
                "previsao_termino": data(g("previsao_termino")), "situacao": texto(g("situacao")),
                "motivo_atraso": texto(g("motivo_atraso")), "ano_per": texto(g("ano_per")),
                "observacao": texto(g("observacao")), "previsto": {}, "executado": {},
                "anos_ant_previsto": None, "anos_ant_executado": None,
            }
            obras.append(atual)
        else:  # linha executado: completa campos da obra que só vieram nela
            for k in ("id_sigicor", "item_per", "descricao", "municipio", "uf", "rodovia"):
                if not atual[k]:
                    atual[k] = texto(g(k))
        atual[f"anos_ant_{tipo}"] = numero(g("anos_anteriores"))
        for j, (m, bloco) in meses.items():
            if j < len(row):
                v = numero(row[j])
                if v is not None or row[j] == 0:
                    atual[tipo][(m, bloco)] = v

    # eventograma (Rio Minas, 4º ano): linha "Eventograma" = obra; linhas seguintes com peso numérico = etapas dela
    pai = None
    for o in obras:
        ev = norm(o["evento"] or "")
        peso = numero(o["evento"]) if ev and not re.search(r"[a-z]", ev) else None
        if peso is not None and pai is not None:
            o.update(nivel="etapa", etapa=o["descricao"], peso_etapa=peso, pai=pai)
        else:
            o.update(nivel="obra", etapa=None, peso_etapa=None, pai=None)
            pai = o if ev == "eventograma" else None

    # mês que aparece em mais de um bloco: vale o 1º valor preenchido; conta conflitos (valores diferentes)
    conflitos = 0
    for o in obras:
        for tipo in ("previsto", "executado"):
            por_mes = {}
            for (m, b) in sorted(o[tipo], key=lambda x: x[1]):
                por_mes.setdefault(m, []).append(o[tipo][(m, b)])
            o[f"{tipo}_mes"] = {}
            for m, vs in por_mes.items():
                preenchidos = [v for v in vs if v is not None]
                o[f"{tipo}_mes"][m] = preenchidos[0] if preenchidos else None
                if len({round(v, 6) for v in preenchidos}) > 1:
                    conflitos += 1

    # "Atualizado - dd/mm/aaaa - Referência ..." no topo (ECO050, Cerrado)
    atualizado, texto_atual = None, None
    for r in topo[:hdr]:
        for c in r:
            m = isinstance(c, str) and re.search(r"atualizado\D{0,5}(\d{2})/(\d{2})/(20\d{2})", c, re.I)
            if m:
                atualizado = datetime(int(m.group(3)), int(m.group(2)), int(m.group(1)))
                texto_atual = texto(c)
    aba = ws.title
    wb.close()
    repetidos = len({m for m, _ in meses.values()}) < len(meses)
    return {"aba": aba, "linha_cabecalho": hdr + 1, "obras": obras, "meses": meses, "atualizado": atualizado,
            "texto_atualizacao": texto_atual, "blocos": rotulos, "meses_repetidos": repetidos,
            "conflitos_blocos": conflitos, "colunas": sorted(col)}


# ----------------------------------------------------------------------------- identidade da obra

def ev_norm(ev) -> str:
    """'Etapa única' = obra sem divisão em eventos (equivale ao modelo antigo, que não tinha a coluna)."""
    e = norm(ev or "") if ev == ev else ""
    return "" if e in ("", "etapa unica", "eventograma") else e


def chave_base(o: dict) -> str:
    desc = re.sub(r"[^a-z0-9+]+", " ", norm(o["descricao"] or "")).strip()
    item = re.sub(r"[^a-z0-9]+", ".", norm(o["item_per"] or "")).strip(".")
    return f"{item}|{desc}|{ev_norm(o['evento'])}"


def indicadores(o: dict, ref: str, blocos_meses: dict, rotulos: dict) -> dict:
    """% acumulados até o mês de referência.
    Executado: avanço real, soma "anos anteriores" + incrementos mensais até ref.
    Previsto: cada ano de concessão traz um plano NOVO (reprograma o que faltou do ano anterior), então somar os
    previstos de anos diferentes passa de 100%. O previsto acumulado é: executado real até o início do ano de
    concessão de ref + previsto desse ano até ref. 'meta_fim_ano' = executado até o início do ano + previsto do ano todo."""
    vazio = {"ano_concessao_ref": None, "pct_previsto_acum_ate_ref": None, "pct_executado_acum_ate_ref": None,
             "pct_previsto_no_mes_ref": None, "pct_executado_no_mes_ref": None, "pct_meta_fim_ano_concessao": None}
    if not ref:
        return vazio
    exe = o["executado_mes"]
    bloco = next((b for b, ms in sorted(blocos_meses.items()) if ref in ms), None)
    if bloco is None:  # ref fora da tabela: usa o último bloco que começa antes de ref
        bloco = max((b for b, ms in blocos_meses.items() if ms and ms[0] <= ref), default=None)
    exec_acum = (o["anos_ant_executado"] or 0) + sum(v or 0 for m, v in exe.items() if m <= ref)
    if bloco is None:
        return dict(vazio, pct_executado_acum_ate_ref=round(exec_acum, 4), pct_executado_no_mes_ref=exe.get(ref))
    inicio = blocos_meses[bloco][0]
    exec_ate_inicio = (o["anos_ant_executado"] or 0) + sum(v or 0 for m, v in exe.items() if m < inicio)
    prev_ano = {m: v for (m, b), v in o["previsto"].items() if b == bloco}
    return {
        "ano_concessao_ref": rotulos.get(bloco),
        "pct_previsto_acum_ate_ref": round(exec_ate_inicio + sum(v or 0 for m, v in prev_ano.items() if m <= ref), 4),
        "pct_executado_acum_ate_ref": round(exec_acum, 4),
        "pct_previsto_no_mes_ref": prev_ano.get(ref),
        "pct_executado_no_mes_ref": exe.get(ref),
        "pct_meta_fim_ano_concessao": round(exec_ate_inicio + sum(v or 0 for v in prev_ano.values()), 4),
    }


def vincular_obras(ob: pd.DataFrame) -> dict:
    """Dá um obra_id estável para cada (snapshot, chave), percorrendo os snapshots em ordem cronológica.
    A concessionária reescreve descrições e às vezes troca o código do item PER entre meses, então a ligação
    tenta, nesta ordem, contra obras já conhecidas que ainda não apareceram no snapshot:
      1. 'exata'               mesma chave (item PER + descrição normalizada) já vista antes;
      2. 'km'                  mesmo item PER, km inicial e final iguais (±50 m); entre várias, a de descrição mais parecida;
      3. 'similaridade'        mesmo item PER, descrição ≥ 85% parecida;
      4. 'item_per_alterado'   outro item PER, descrição ≥ 93% parecida e km compatível (ou ambos sem km);
      senão                    'nova' obra.
    Devolve {(snapshot, chave): (obra_id, metodo)}."""
    def km_ok(a, r):
        def igual(x, y):
            vazio_x, vazio_y = x != x or x is None, y != y or y is None
            return (vazio_x and vazio_y) or (not vazio_x and not vazio_y and abs(x - y) <= 0.05)
        return igual(a.km_inicial, r.km_inicial) and igual(a.km_final, r.km_final)

    def parecido(a, r):
        return SequenceMatcher(None, norm(a.descricao or ""), norm(r.descricao or "")).ratio()

    res = {}
    ob = ob.rename(columns={"_chave": "chave"})  # itertuples não aceita nome começando com "_"
    for conc, g in ob.sort_values(["mes_referencia", "snapshot"]).groupby("concessao", sort=False):
        sigla = SIGLA.get(conc, conc)
        por_chave, registro, n = {}, {}, 0  # chave -> id ; id -> atributos da última aparição
        for snap, h in g.groupby("snapshot", sort=False):
            usados, pendentes = set(), []
            for r in h.itertuples(index=False):
                oid = por_chave.get(r.chave)
                if oid and oid not in usados:
                    usados.add(oid)
                    res[(snap, r.chave)] = (oid, "exata")
                    registro[oid] = r
                else:
                    pendentes.append(r)
            for r in pendentes:
                livres = [(oid, a) for oid, a in registro.items() if oid not in usados]
                mesmo_item = [(oid, a) for oid, a in livres if (a.item_per or "") == (r.item_per or "")]
                escolha, metodo = None, None
                tem_km = r.km_inicial == r.km_inicial and r.km_inicial is not None
                if tem_km:
                    por_km = [(parecido(a, r), oid) for oid, a in mesmo_item if km_ok(a, r)]
                    if por_km:
                        escolha, metodo = max(por_km)[1], "km"
                if not escolha:
                    melhor = max(((parecido(a, r), oid) for oid, a in mesmo_item), default=(0, None))
                    if melhor[0] >= 0.85:
                        escolha, metodo = melhor[1], "similaridade"
                if not escolha:
                    outros = [(parecido(a, r), oid) for oid, a in livres
                              if (a.item_per or "") != (r.item_per or "") and km_ok(a, r)]
                    melhor = max(outros, default=(0, None))
                    if melhor[0] >= 0.93:
                        escolha, metodo = melhor[1], "item_per_alterado"
                if not escolha:
                    n += 1
                    escolha, metodo = f"{sigla}-{n:04d}", "nova"
                usados.add(escolha)
                por_chave[r.chave] = escolha
                registro[escolha] = r
                res[(snap, r.chave)] = (escolha, metodo)
    return res


def pct(x) -> str:
    return "?" if x is None or x != x else f"{x * 100:.1f}%".replace(".", ",")


def marcos_por_obra(pr: pd.DataFrame, mud: pd.DataFrame, se: pd.DataFrame) -> pd.DataFrame:
    """Uma linha por obra com os marcos, em duas visões:
    - "relatado": o que a concessionária dizia em cada mês (obras_por_snapshot, snapshots principais);
    - "último arquivo": o que a planilha mais recente da obra diz sobre os meses passados (serie_mensal).
    Quando as duas divergem, houve revisão retroativa. Limiar de 0,05 p.p. para ignorar arredondamento."""
    EPS = 0.0005
    pr = pr.sort_values(["obra_id", "mes_referencia"])
    refs_conc = {c: sorted(g["mes_referencia"].unique()) for c, g in pr.groupby("concessao")}
    ex = se[(se["nivel"] == "obra") & (se["tipo"] == "executado")]
    primeiro_exec = ex[ex["pct"] > EPS].groupby(["snapshot", "obra_id"])["mes"].min()
    primeiro_mes_tabela = se.groupby("snapshot")["mes"].min()
    mud = mud if len(mud) else pd.DataFrame(columns=["obra_id", "para_mes", "executado_passado_alterado",
                                                     "previsto_passado_alterado", "situacao_alterada"])
    mud_por_obra = {oid: g for oid, g in mud.groupby("obra_id")}

    linhas = []
    for oid, g in pr.groupby("obra_id", sort=False):
        refs = g["mes_referencia"].tolist()
        exe = g["pct_executado_acum_ate_ref"].tolist()
        sit = g["situacao"].fillna("").str.upper().tolist()
        ult = g.iloc[-1]
        conc = ult["concessao"]

        i_av = next((i for i, v in enumerate(exe) if v and v > EPS), None)
        i_100 = next((i for i, v in enumerate(exe) if v and v >= 0.995), None)
        i_conc = next((i for i, s in enumerate(sit) if s.startswith("CONCLU")), None)
        variacoes = [(refs[i], (exe[i] or 0) - (exe[i - 1] or 0)) for i in range(1, len(exe))
                     if abs((exe[i] or 0) - (exe[i - 1] or 0)) > 0.0001]
        # chegou a 100% e depois caiu (replanejamento, obra dividida ou ligação de identidade errada)
        i_caiu = next((i for i in range(i_100 + 1, len(exe)) if (exe[i] or 0) < 0.995), None) if i_100 is not None else None

        # início segundo o último arquivo da obra
        snap = ult["snapshot"]
        if (ult["pct_executado_anos_anteriores"] or 0) > EPS:
            inicio_atual = f"antes de {primeiro_mes_tabela.get(snap)}"
        else:
            inicio_atual = primeiro_exec.get((snap, oid))
        # revisado: o último arquivo põe o início depois do 1º relato, ou antes num mês em que a obra
        # foi relatada com 0%
        revisado = False
        if i_av is not None and inicio_atual and not inicio_atual.startswith("antes"):
            revisado = inicio_atual > refs[i_av] or any(inicio_atual <= r < refs[i_av] for r in refs[:i_av])

        m = mud_por_obra.get(oid, mud.iloc[0:0])
        rev_exec = m[m["executado_passado_alterado"] == True]["para_mes"].tolist()
        rep_prev = m[m["previsto_passado_alterado"] == True]["para_mes"].tolist()
        conc_refs = [r for r in refs_conc[conc] if refs[0] <= r <= refs[-1]]
        ausentes = [r for r in conc_refs if r not in set(refs)]
        vinculos = sorted(set(g["vinculo"]) - {"exata", "nova"})
        presente = refs[-1] == refs_conc[conc][-1]

        partes = [f"Surgiu em {refs[0]} ({sit[0] or 'sem situação'}, {pct(exe[0])} executado)."]
        if i_av is None:
            partes.append("Nenhum avanço relatado até hoje.")
        elif i_av == 0 and exe[0] > EPS:
            partes.append("Já tinha execução quando apareceu pela 1ª vez.")
        else:
            partes.append(f"1º avanço relatado em {refs[i_av]} ({pct(exe[i_av])}).")
        if revisado:
            partes.append(f"Na planilha mais recente o início da execução está em {inicio_atual} (revisão retroativa).")
        if rev_exec:
            partes.append(f"Executado de meses passados revisado {len(rev_exec)}x (última em {rev_exec[-1]}).")
        if rep_prev:
            partes.append(f"Previsto de meses passados reescrito {len(rep_prev)}x (última em {rep_prev[-1]}).")
        if i_100 is not None:
            partes.append(f"Atingiu 100% em {refs[i_100]}.")
        anos_conc = g["ano_concessao_ref"].tolist()
        novo_ciclo = i_caiu is not None and anos_conc[i_caiu] != anos_conc[i_caiu - 1]
        if i_caiu is not None:
            partes.append(f"Depois voltou a {pct(exe[i_caiu])} em {refs[i_caiu]}"
                          + (" (novo ano de concessão: serviço recorrente reiniciado)." if novo_ciclo else "."))
        if i_conc is not None:
            partes.append(f"Situação 'CONCLUÍDA' a partir de {refs[i_conc]}.")
        elif i_100 is not None:
            partes.append(f"Situação nunca foi 'CONCLUÍDA' (última: {ult['situacao']}).")
        if variacoes:
            r, d = variacoes[-1]
            pp = f"{d * 100:+.2f}".replace(".", ",")
            partes.append(f"Última variação do executado em {r} ({pp} p.p.).")
        if ausentes:
            partes.append(f"Ausente em {len(ausentes)} mês(es) publicados: {', '.join(ausentes[:6])}.")
        if not presente:
            partes.append(f"Saiu da tabela: última aparição em {refs[-1]}.")
        if vinculos:
            partes.append(f"Identidade ligada entre meses por aproximação ({', '.join(vinculos)}): conferir.")
        partes.append(f"Último dado ({refs[-1]}): {pct(exe[-1])} executado, {pct(ult['pct_previsto_acum_ate_ref'])} previsto,"
                      f" {ult['situacao'] or 'sem situação'}.")

        linhas.append({
            "id_empreendimento": ult["id_empreendimento"], "concessao": conc, "obra_id": oid,
            "item_per": ult["item_per"], "descricao": ult["descricao"], "rodovia": ult["rodovia"],
            "km_inicial": ult["km_inicial"], "km_final": ult["km_final"],
            "surgiu_em": refs[0], "situacao_ao_surgir": g["situacao"].iloc[0], "pct_exec_ao_surgir": exe[0],
            "ja_iniciada_ao_surgir": bool(exe[0] and exe[0] > EPS),
            "primeiro_avanco_relatado_em": refs[i_av] if i_av is not None else None,
            "pct_no_primeiro_avanco": exe[i_av] if i_av is not None else None,
            "inicio_execucao_ultimo_arquivo": inicio_atual,
            "inicio_revisado_retroativamente": revisado,
            "atingiu_100_em": refs[i_100] if i_100 is not None else None,
            "situacao_concluida_em": refs[i_conc] if i_conc is not None else None,
            "caiu_abaixo_100_em": refs[i_caiu] if i_caiu is not None else None,
            "queda_na_virada_do_ano_concessao": novo_ciclo,
            "ultima_variacao_exec_em": variacoes[-1][0] if variacoes else None,
            "ultima_variacao_exec_pp": round(variacoes[-1][1] * 100, 2) if variacoes else None,
            "meses_com_variacao_exec": len(variacoes),
            "revisoes_exec_retroativas": len(rev_exec),
            "ultima_revisao_exec_retroativa_em": rev_exec[-1] if rev_exec else None,
            "reprogramacoes_previsto_passado": len(rep_prev),
            "mudancas_situacao": int((m["situacao_alterada"] == True).sum()),
            "ultima_ref": refs[-1], "situacao_ultima": ult["situacao"], "pct_exec_ultimo": exe[-1],
            "pct_prev_ultimo": ult["pct_previsto_acum_ate_ref"], "presente_no_ultimo_snapshot": presente,
            "meses_ausente": ", ".join(ausentes), "vinculo_aproximado": ", ".join(vinculos),
            "alerta_acima_100": g["alerta"].notna().any(),
            "historico": " ".join(partes),
        })
    return pd.DataFrame(linhas).sort_values(["id_empreendimento", "obra_id"])


def main(argv):
    alvo = [s for s in EMPS if not argv or s in argv]
    SAIDA.mkdir(parents=True, exist_ok=True)
    snaps, linhas_obra, linhas_etapa, serie = [], [], [], []

    for slug in alvo:
        arquivos = listar_planilhas(slug)
        print(f"[{slug}] {len(arquivos)} planilhas")
        for p in arquivos:
            try:
                lido = ler_planilha(p)
            except Exception as e:
                print(f"    ERRO {p.name}: {e}")
                continue
            if not lido or not lido["obras"]:
                continue
            rel = p.relative_to(BRUTOS / slug)
            ano_pasta = int(rel.parts[0])
            origem = rel.parts[2] if rel.parts[1] == "extraido" else ""
            pub, mes_nome, _ = datas_do_nome(p.stem)
            if origem:  # arquivo dentro de ZIP: completa pelo nome do ZIP
                pub2, mes2, _ = datas_do_nome(origem)
                pub, mes_nome = pub or pub2, mes_nome or mes2
            pub = pub or lido["atualizado"]

            # o que os próprios dados indicam (sem passar da publicação):
            #  - último mês com executado PREENCHIDO (inclui zero): bom quando o mês futuro fica em branco (ECO050)
            #  - último mês com executado > 0: necessário quando o ano todo vem preenchido com zeros (Rio Minas, 4º ano)
            limite = pub.strftime("%Y-%m") if pub else "2100-01"
            preench = sorted({m for o in lido["obras"] for m, v in o["executado_mes"].items() if v is not None and m < limite})
            avanco = sorted({m for o in lido["obras"] for m, v in o["executado_mes"].items() if v and v > 0 and m < limite})
            ult_preench = preench[-1] if preench else None
            ult_avanco = avanco[-1] if avanco else None
            nmes = lambda c: int(c[:4]) * 12 + int(c[5:7])
            zeros_adiante = bool(ult_preench and ult_avanco and nmes(ult_preench) - nmes(ult_avanco) > 2)
            mes_dados = ult_avanco if zeros_adiante else ult_preench

            alerta = []
            if mes_nome:
                ref, fonte = resolver_mes(mes_nome, pub, ano_pasta, mes_dados), "nome do arquivo"
            elif pub:
                ref = (pub.replace(day=1) - pd.DateOffset(months=1)).strftime("%Y-%m")
                fonte = "data de publicação - 1 mês"
            elif mes_dados:
                ref = mes_dados
                fonte = "dados: último mês com avanço > 0" if zeros_adiante else "dados: último mês com executado preenchido"
                alerta.append("sem mês no nome nem data de publicação: mês deduzido dos dados")
            else:
                ref, fonte = None, ""
            if ref and mes_dados and abs(nmes(ref) - nmes(mes_dados)) > 1:
                alerta.append(f"dados sugerem {mes_dados}")
            if lido["conflitos_blocos"]:
                alerta.append(f"{lido['conflitos_blocos']} valores divergentes entre blocos repetidos (usado o 1º)")

            blocos_meses = {}
            for m, b in lido["meses"].values():
                blocos_meses.setdefault(b, []).append(m)
            blocos_meses = {b: sorted(ms) for b, ms in blocos_meses.items()}

            n_snap = sum(1 for s in snaps if s["concessao"] == slug) + 1
            snap_id = f"{SIGLA.get(slug, slug)}-{ref or 'sem-mes'}-{n_snap:02d}"
            snaps.append({
                "snapshot": snap_id, "id_empreendimento": ID_EMP[slug], "concessao": slug,
                "mes_referencia": ref, "fonte_mes_ref": fonte,
                "data_publicacao": pub.date().isoformat() if pub else None,
                "texto_atualizacao": lido["texto_atualizacao"],
                "ult_mes_exec_preenchido": ult_preench, "ult_mes_com_avanco": ult_avanco,
                "arquivo": p.name, "zip": origem, "pasta_ano": ano_pasta, "aba": lido["aba"],
                "linha_cabecalho": lido["linha_cabecalho"], "blocos": " | ".join(lido["blocos"].values()),
                "n_obras": sum(1 for o in lido["obras"] if o["nivel"] == "obra"),
                "n_etapas": sum(1 for o in lido["obras"] if o["nivel"] == "etapa"), "alerta": "; ".join(alerta),
            })

            vistos = {}
            for o in lido["obras"]:
                prev, exe = o["previsto_mes"], o["executado_mes"]
                if o["nivel"] == "etapa":
                    o["_chave"] = o["pai"]["_chave"]
                    linhas_etapa.append({
                        "snapshot": snap_id, "concessao": slug, "mes_referencia": ref, "_chave": o["_chave"],
                        "etapa": o["etapa"], "peso_etapa": o["peso_etapa"], "situacao": o["situacao"],
                        **indicadores(o, ref, blocos_meses, lido["blocos"]),
                    })
                    for tipo in ("previsto", "executado"):
                        for (m, b), v in o[tipo].items():
                            serie.append({"snapshot": snap_id, "concessao": slug, "mes_referencia": ref, "_chave": o["_chave"],
                                          "nivel": "etapa", "etapa": o["etapa"], "bloco": lido["blocos"].get(b, b),
                                          "mes": m, "tipo": tipo, "pct": v})
                    continue
                base = chave_base(o)
                vistos[base] = vistos.get(base, 0) + 1
                o["_chave"] = base if vistos[base] == 1 else f"{base}|#{vistos[base]}"
                linhas_obra.append({
                    "snapshot": snap_id, "id_empreendimento": ID_EMP[slug], "concessao": slug, "mes_referencia": ref,
                    "data_publicacao": snaps[-1]["data_publicacao"], "_chave": o["_chave"],
                    **{k: o[k] for k in ("id_sigicor", "item_per", "descricao", "evento", "frente", "municipio", "uf",
                                         "rodovia", "km_inicial", "km_final", "data_inicio", "previsao_termino")},
                    "situacao": o["situacao"],
                    **indicadores(o, ref, blocos_meses, lido["blocos"]),
                    "pct_executado_anos_anteriores": o["anos_ant_executado"],
                    "motivo_atraso": o["motivo_atraso"], "observacao": o["observacao"], "ano_per": o["ano_per"],
                })
                for tipo in ("previsto", "executado"):
                    for (m, b), v in o[tipo].items():
                        serie.append({"snapshot": snap_id, "concessao": slug, "mes_referencia": ref, "_chave": o["_chave"],
                                      "nivel": "obra", "etapa": None, "bloco": lido["blocos"].get(b, b),
                                      "mes": m, "tipo": tipo, "pct": v})
            print(f"    {ref or '???????'} {fonte[:24]:24s} obras={snaps[-1]['n_obras']:4d} etapas={snaps[-1]['n_etapas']:4d}  {p.name[:62]}"
                  + (f"   ! {snaps[-1]['alerta']}" if snaps[-1]["alerta"] else ""))

    s = pd.DataFrame(snaps)
    ob = pd.DataFrame(linhas_obra)
    se = pd.DataFrame(serie)

    # valores que passam de 100%: na fonte, alguns meses vêm como % ACUMULADO em vez de incremento mensal
    acima = (ob["pct_previsto_acum_ate_ref"] > 1.05) | (ob["pct_executado_acum_ate_ref"] > 1.05)
    ob["alerta"] = acima.map({True: "soma > 100%: a fonte provavelmente lançou % acumulado em vez de mensal", False: None})

    # vários arquivos para o mesmo mês: vale o publicado por último (depois, o de mais linhas)
    s["_ord_pub"] = s["data_publicacao"].fillna("")
    s = s.sort_values(["concessao", "mes_referencia", "_ord_pub", "n_obras", "arquivo"])
    s["principal"] = ~s.duplicated(["concessao", "mes_referencia"], keep="last") & s["mes_referencia"].notna()
    s = s.drop(columns="_ord_pub").sort_values(["concessao", "mes_referencia", "arquivo"])

    # id estável da obra (ver vincular_obras)
    ob = ob.merge(s[["snapshot", "principal"]], on="snapshot")
    ob = ob.sort_values(["concessao", "mes_referencia", "snapshot"])
    vinc = vincular_obras(ob)
    ob.insert(1, "obra_id", [vinc[(sn, k)][0] for sn, k in zip(ob["snapshot"], ob["_chave"])])
    ob.insert(2, "vinculo", [vinc[(sn, k)][1] for sn, k in zip(ob["snapshot"], ob["_chave"])])
    se.insert(1, "obra_id", [vinc[(sn, k)][0] for sn, k in zip(se["snapshot"], se["_chave"])])
    et = pd.DataFrame(linhas_etapa)
    if len(et):
        et.insert(1, "obra_id", [vinc[(sn, k)][0] for sn, k in zip(et["snapshot"], et["_chave"])])
        et = et.drop(columns="_chave")

    # catálogo: dados da obra na última aparição (snapshots principais)
    pr = ob[ob["principal"]]
    ult_snap = pr.groupby("concessao")["mes_referencia"].max()
    cat = (pr.sort_values("mes_referencia").groupby("obra_id")
           .agg(id_empreendimento=("id_empreendimento", "last"), concessao=("concessao", "last"), item_per=("item_per", "last"), descricao=("descricao", "last"),
                evento=("evento", "last"), id_sigicor=("id_sigicor", "last"), frente=("frente", "last"),
                municipio=("municipio", "last"), uf=("uf", "last"), rodovia=("rodovia", "last"),
                km_inicial=("km_inicial", "last"), km_final=("km_final", "last"),
                primeira_ref=("mes_referencia", "min"), ultima_ref=("mes_referencia", "max"),
                n_snapshots=("snapshot", "nunique"), situacao_ultima=("situacao", "last"),
                pct_executado_ultimo=("pct_executado_acum_ate_ref", "last")).reset_index())
    cat["presente_no_ultimo_snapshot"] = cat["ultima_ref"] == cat["concessao"].map(ult_snap)

    # lista enxuta para comparar com outras bases (ex.: PELT): nome atual + nomes/códigos que a obra já teve
    def outros(col):
        return lambda s: " | ".join(dict.fromkeys(x for x in s[::-1].dropna().astype(str)))
    hist = (pr.sort_values("mes_referencia").groupby("obra_id")
            .agg(descricoes_anteriores=("descricao", outros("descricao")), itens_per=("item_per", outros("item_per"))))
    lista = cat.merge(hist, on="obra_id")
    lista["empreendimento"] = lista["concessao"].map(lambda s: EMPS[s]["nome"])
    lista["descricoes_anteriores"] = [
        " | ".join(d for d in todas.split(" | ") if d != atual) or None
        for todas, atual in zip(lista["descricoes_anteriores"], lista["descricao"].astype(str))]
    lista = lista[["id_empreendimento", "empreendimento", "obra_id", "item_per", "id_sigicor", "descricao", "rodovia",
                   "km_inicial", "km_final", "primeira_ref", "ultima_ref", "presente_no_ultimo_snapshot",
                   "itens_per", "descricoes_anteriores"]].sort_values(["id_empreendimento", "obra_id"])

    # previsto mês a mês de cada obra em cada snapshot (1º valor preenchido entre blocos)
    prev_mes = (se[(se["nivel"] == "obra") & (se["tipo"] == "previsto")]
                .groupby(["snapshot", "obra_id", "mes"], sort=False)["pct"].first().sort_index())
    exec_mes = (se[(se["nivel"] == "obra") & (se["tipo"] == "executado")]
                .groupby(["snapshot", "obra_id", "mes"], sort=False)["pct"].first().sort_index())

    def diff_meses(serie, sa, sb, oid, ate=None, depois=None):
        try:
            a, b = serie.loc[(sa, oid)], serie.loc[(sb, oid)]
        except KeyError:
            return []
        j = pd.concat([a, b], axis=1, keys=["a", "b"]).dropna(how="all")
        if ate:
            j = j[j.index <= ate]
        if depois:
            j = j[j.index > depois]
        return sorted(j.index[(j["a"].fillna(0) - j["b"].fillna(0)).abs() > 0.005])

    evo, mud = [], []
    for conc, g in pr.groupby("concessao"):
        ant, ant_snap, ant_ref = None, None, None
        for r, h in g.sort_values("mes_referencia").groupby("mes_referencia"):
            atual = h.set_index("obra_id")
            snap = h["snapshot"].iloc[0]
            linha = {"concessao": conc, "mes_referencia": r, "snapshot": snap, "obras": len(atual),
                     "pct_exec_medio": round(atual["pct_executado_acum_ate_ref"].mean(), 4)}
            for st, n in atual["situacao"].fillna("(vazio)").str.upper().value_counts().items():
                linha[f"situacao: {st}"] = n
            if ant is not None:
                comuns = atual.index.intersection(ant.index)
                linha["obras_novas"] = len(set(atual.index) - set(ant.index))
                linha["obras_que_sairam"] = len(set(ant.index) - set(atual.index))
                c = dict(passado=0, futuro=0, exec_revisado=0, exec_reduzido=0, situacao=0, termino=0, descricao=0, item=0)
                for oid in comuns:
                    a, b = ant.loc[oid], atual.loc[oid]
                    m_pass = diff_meses(prev_mes, ant_snap, snap, oid, ate=ant_ref)
                    m_fut = diff_meses(prev_mes, ant_snap, snap, oid, depois=ant_ref)
                    m_exec = diff_meses(exec_mes, ant_snap, snap, oid, ate=ant_ref)
                    d_exec = (b["pct_executado_acum_ate_ref"] or 0) - (a["pct_executado_acum_ate_ref"] or 0)
                    flags = {
                        "previsto_passado_alterado": bool(m_pass), "previsto_futuro_alterado": bool(m_fut),
                        "executado_passado_alterado": bool(m_exec), "executado_acum_reduzido": d_exec < -0.005,
                        "situacao_alterada": (a["situacao"] or "") != (b["situacao"] or ""),
                        "previsao_termino_alterada": (a["previsao_termino"] or "") != (b["previsao_termino"] or ""),
                        "descricao_alterada": (a["descricao"] or "") != (b["descricao"] or ""),
                        "item_per_alterado": (a["item_per"] or "") != (b["item_per"] or ""),
                    }
                    c["passado"] += flags["previsto_passado_alterado"]
                    c["futuro"] += flags["previsto_futuro_alterado"]
                    c["exec_revisado"] += flags["executado_passado_alterado"]
                    c["exec_reduzido"] += flags["executado_acum_reduzido"]
                    c["situacao"] += flags["situacao_alterada"]
                    c["termino"] += flags["previsao_termino_alterada"]
                    c["descricao"] += flags["descricao_alterada"]
                    c["item"] += flags["item_per_alterado"]
                    if any(flags.values()) or abs(d_exec) > 0.005:
                        mud.append({
                            "concessao": conc, "obra_id": oid, "de_mes": ant_ref, "para_mes": r,
                            "item_per": b["item_per"], "descricao": b["descricao"],
                            "situacao_de": a["situacao"], "situacao_para": b["situacao"],
                            "pct_exec_de": a["pct_executado_acum_ate_ref"], "pct_exec_para": b["pct_executado_acum_ate_ref"],
                            "delta_exec": round(d_exec, 4),
                            "pct_prev_de": a["pct_previsto_acum_ate_ref"], "pct_prev_para": b["pct_previsto_acum_ate_ref"],
                            **flags,
                            "meses_previsto_passado_alterados": ", ".join(m_pass[:12]),
                            "meses_executado_passado_alterados": ", ".join(m_exec[:12]),
                            "previsao_termino_de": a["previsao_termino"], "previsao_termino_para": b["previsao_termino"],
                            "descricao_de": a["descricao"] if flags["descricao_alterada"] else None,
                            "item_per_de": a["item_per"] if flags["item_per_alterado"] else None,
                        })
                linha.update({
                    "obras_com_previsto_passado_reescrito": c["passado"], "obras_com_previsto_futuro_alterado": c["futuro"],
                    "obras_com_executado_passado_revisado": c["exec_revisado"],
                    "obras_com_executado_acum_reduzido": c["exec_reduzido"],
                    "obras_com_situacao_alterada": c["situacao"], "obras_com_previsao_termino_alterada": c["termino"],
                    "obras_com_descricao_alterada": c["descricao"], "obras_com_item_per_alterado": c["item"],
                })
            evo.append(linha)
            ant, ant_snap, ant_ref = atual, snap, r
    evo = pd.DataFrame(evo)
    mud = pd.DataFrame(mud)
    marcos = marcos_por_obra(pr, mud, se)

    ob = ob.drop(columns="_chave")
    se = se.drop(columns="_chave")
    s.to_csv(SAIDA / "snapshots.csv", index=False, sep=";", encoding="utf-8-sig")
    cat.to_csv(SAIDA / "obras.csv", index=False, sep=";", encoding="utf-8-sig")
    marcos.to_csv(SAIDA / "marcos_por_obra.csv", index=False, sep=";", encoding="utf-8-sig")
    lista.to_csv(SAIDA / "lista_obras.csv", index=False, sep=";", encoding="utf-8-sig")
    lista.to_excel(SAIDA / "lista_obras.xlsx", index=False)
    ob.to_csv(SAIDA / "obras_por_snapshot.csv", index=False, sep=";", encoding="utf-8-sig")
    se.to_csv(SAIDA / "serie_mensal.csv", index=False, sep=";", encoding="utf-8-sig")
    et.to_csv(SAIDA / "etapas_por_snapshot.csv", index=False, sep=";", encoding="utf-8-sig")
    evo.to_csv(SAIDA / "evolucao.csv", index=False, sep=";", encoding="utf-8-sig")
    mud.to_csv(SAIDA / "mudancas_por_obra.csv", index=False, sep=";", encoding="utf-8-sig")
    with pd.ExcelWriter(SAIDA / "consolidado.xlsx") as xw:
        s.to_excel(xw, "snapshots", index=False)
        cat.to_excel(xw, "obras", index=False)
        marcos.to_excel(xw, "marcos_por_obra", index=False)
        ob.to_excel(xw, "obras_por_snapshot", index=False)
        evo.to_excel(xw, "evolucao", index=False)
        mud.to_excel(xw, "mudancas_por_obra", index=False)

    print(f"\n{len(s)} snapshots ({int(s['principal'].sum())} principais) | {len(cat)} obras | "
          f"{len(ob)} linhas obra x snapshot | {len(se)} pontos na série mensal -> {SAIDA}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
